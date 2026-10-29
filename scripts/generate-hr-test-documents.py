"""生成虚构人事制度的 DOCX 和带中文文字层的 PDF；不调用模型或联网。"""

from pathlib import Path
import argparse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "examples" / "hr-test-data"
COMPANY = "澄远信息技术有限公司"
NOTICE = "本公司及本文全部规则均为虚构测试数据，仅用于系统测试，不作为实际劳动管理依据。"

RECRUITMENT = [
    ("一、适用范围与招聘流程", [
        "本制度适用于公司正式岗位招聘及员工人才推荐。实习岗位、外包岗位以及未在招聘平台发布的岗位，不参与本制度的推荐奖金计划。",
        "用人部门提出招聘需求，由部门负责人确认编制和预算，人力资源部发布岗位。人力资源部完成简历初筛，直属用人主管组织岗位面试，必要时增加专业评估。录用通知由人力资源部统一发出，员工不得擅自承诺录用、薪酬或奖金。",
        "候选人个人信息仅用于招聘和必要的录用审查。推荐人须事先取得候选人同意，不得未经同意在公开群聊中传播简历。",
    ]),
    ("二、推荐奖励标准", [
        "岗位类别以提交推荐当日招聘平台标注为准。普通岗位推荐奖励总额为5000元/人；关键技术岗位推荐奖励总额为8000元/人；管理岗位推荐奖励总额为12000元/人。上述金额均为税前金额，不与岗位月薪挂钩。",
        "同一候选人只计一次有效推荐，不能同时按普通岗位和关键技术岗位重复领取。候选人最终入职岗位发生变化时，由人力资源部在录用确认时说明适用类别并留存双方确认记录。",
    ]),
    ("三、分期发放与时间条件", [
        "推荐奖金分两次发放。第一次为总额的30%，在候选人正式入职满30个自然日、仍在职且推荐记录审核通过后，随下一次月度工资发放。第二次为总额的70%，在候选人通过试用期转正、仍在职且审核通过后，随下一次月度工资发放。",
        "例如：普通岗位总额5000元，第一次1500元，第二次3500元；关键技术岗位总额8000元，第一次2400元，第二次5600元；管理岗位总额12000元，第一次3600元，第二次8400元。",
        "试用期长短以候选人的劳动合同为准，不将入职满30天等同于转正。已经完成发放条件并确认的奖金按已确认结果处理；尚未满足条件的后续奖励，按实际入职和转正情况审核。本制度不设置统一的试用期天数。",
    ]),
    ("四、有效推荐与重复简历", [
        "员工应通过招聘平台的“内部推荐”入口提交候选人姓名、目标岗位、联系方式和简历。人力资源部原则上在收到完整资料后5个工作日内反馈是否为有效推荐。通过私人聊天转发简历但未登记的，不视为已完成有效推荐。",
        "候选人在有效推荐提交前90个自然日内，已进入公司人才库且已被招聘人员有效接触的，不重复计为奖励推荐。有效接触包括已完成的电话沟通、面试邀约或面试；单纯保存一份从未联系的简历不等同于有效接触。",
        "多人推荐同一候选人时，以招聘平台中最早一条经核实有效的完整推荐记录为准。不能仅依据截图发送时间决定归属。",
    ]),
    ("五、不适用奖金的人员与情况", [
        "承担该岗位招聘职责的人力资源人员、该岗位的直接录用决策人，不参与该岗位的推荐奖金。其他员工可按规则参加。",
        "候选人由猎头或其他付费招聘渠道已有效引入的，不再重复发放员工推荐奖金。虚构候选人、冒用他人推荐或伪造推荐时间的，取消对应未确认奖励并按合法有效的流程处理。",
        "亲属推荐应如实申报关系，由人力资源部和用人部门评估利益冲突，不得隐瞒关系或绕过正常面试。亲属关系本身不直接决定是否录用。",
    ]),
    ("六、奖金申请与异议", [
        "推荐人达到分期条件后，在招聘平台确认发放申请。人力资源部核查推荐记录、候选人在职及转正情况，财务部核查金额并按适用规则处理税费。税后实发金额取决于实际计税情况，本制度不提供固定税率。",
        "推荐人对有效性、岗位类别或奖金金额存在异议，应在收到结果后10个工作日内提交“人事咨询”工单，附推荐记录编号。人力资源部原则上在2个工作日内首次响应。",
    ]),
]

TRAVEL = [
    ("一、适用范围与出差审批", [
        "本制度适用于员工因公司业务安排发生的中国境内出差，不适用于日常上下班通勤、个人旅行或境外差旅。实习人员因公出行时由业务负责人事先确认适用安排。",
        "出差原则上至少提前3个工作日申请。申请应填写出差目的、城市、起止时间、预计费用以及费用承担部门，由直属主管和部门负责人审批。突发业务确需紧急出行的，应先取得可追溯的审批确认，并于出发后2个工作日内补齐系统申请。",
        "出差行程超过7个自然日时，须增加分管负责人审批。个人自行延长行程产生的住宿、交通和餐费不纳入公司报销。",
    ]),
    ("二、交通与住宿标准", [
        "境内铁路原则上选择高铁或动车二等座、普通列车硬卧；境内航班原则上选择经济舱。升级席位应在订票前取得部门负责人书面审批。市内因公交通凭有效凭证按实际合理费用报销，个人绕行或旅游费用不得混入。",
        "住宿标准按实际住宿城市判断，单位为每人每晚。北京、上海、深圳的住宿上限为600元；其他中国境内城市的住宿上限为450元。住宿限额包含房费和随房费收取的必要税费，不含个人迷你吧、娱乐或消费品费用。",
        "在标准以内按实际发生额报销，不按上限固定发放。超出住宿上限的，须在预订前取得部门负责人和财务负责人书面审批，并在报销时附审批记录。未事先审批的超标费用应提交说明，由公司按有效规则核实处理，不能自行认定全额可报。",
    ]),
    ("三、餐费补助与统一供餐", [
        "境内出差餐费补助标准为每人每天100元，按经批准的实际出差自然日计发。餐费补助不与同一餐次的餐饮发票重复报销。",
        "公司、客户或活动组织方已统一免费提供的餐次，应从当天补助中扣减：早餐按20%、午餐按40%、晚餐按40%扣减。酒店房费已包含早餐的，视为统一提供早餐。未统一供餐的餐次按剩余比例计算。",
        "例如：三天出差期间，客户每天统一提供午餐和晚餐，且酒店不含早餐，则每天补助20元，三天共60元。若同一天三餐均免费提供，则当天餐费补助为0元。",
    ]),
    ("四、报销期限与材料", [
        "员工应在返程后10个工作日内提交报销申请。因客观原因不能按期提交时，应在期限内向直属主管和财务部说明原因及预计提交日期。",
        "报销材料至少包含：已批准的出差审批单、按日期和类别整理的费用明细、符合要求的有效发票、机票或火车票等行程证明。申请超标住宿或升级席位的，还需相应事先审批记录。餐费补助需如实标注统一供餐情况。",
        "报销由直属主管确认业务真实性，部门负责人确认预算，财务部核验票据、标准及是否重复。资料不完整时，财务部应说明需要补充的项目；员工补齐后再进入后续审核。",
    ]),
    ("五、计算示例与边界", [
        "上海出差住3晚，每晚580元：城市住宿上限600元，每晚未超标，住宿费用合计1740元。是否最终报销仍需满足出差审批、真实性和票据要求。",
        "成都出差酒店每晚480元：成都属于其他境内城市，常规上限为450元/人/晚，超出30元/晚，预订前需要部门负责人和财务负责人书面审批。不得把上海标准套用于成都。",
        "本制度未提供境外住宿标准、里程补贴、私车公用补贴或固定审批到账天数。遇到未覆盖事项，应向财务部咨询，不得自动套用本制度的境内标准。",
    ]),
]


def set_word_font(run, size=11, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def create_word():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    doc.core_properties.title = "招聘与人才推荐奖励制度"
    doc.core_properties.author = COMPANY + "（虚构测试）"
    for line, size in [(COMPANY, 12), ("招聘与人才推荐奖励制度", 22)]:
        p = doc.add_paragraph()
        set_word_font(p.add_run(line), size, True, "17554A")
    for line in ["文档编号：HR-003    版本：V1.0", "生效日期：2026年10月1日    制度归口：人力资源部", NOTICE]:
        p = doc.add_paragraph()
        set_word_font(p.add_run(line), 10)
    for title, paragraphs in RECRUITMENT:
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        set_word_font(heading.add_run(title), 14, True, "17554A")
        for text in paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.25
            set_word_font(p.add_run(text))
    footer = section.footer.paragraphs[0]
    set_word_font(footer.add_run("虚构测试资料 | HR-003 | 第 "), 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_word_font(footer.add_run(" 页"), 9)
    doc.save(OUTPUT / "招聘与人才推荐奖励制度.docx")


def create_pdf(font_path):
    # 嵌入本地中文字体，并由 ReportLab 生成 Unicode 映射，保证可复制及被解析器提取。
    pdfmetrics.registerFont(TTFont("Chinese", str(font_path)))
    body = ParagraphStyle("Body", fontName="Chinese", fontSize=10.5, leading=18,
                          wordWrap="CJK", spaceAfter=8, textColor=colors.HexColor("#24322F"))
    heading = ParagraphStyle("Heading", parent=body, fontSize=13, leading=21,
                             spaceBefore=12, textColor=colors.HexColor("#17554A"), keepWithNext=True)
    title = ParagraphStyle("Title", parent=body, fontSize=22, leading=32, spaceAfter=13)
    meta = ParagraphStyle("Meta", parent=body, fontSize=9, leading=15, textColor=colors.HexColor("#53615D"))
    flow = [Paragraph(COMPANY, meta), Spacer(1, 7 * mm), Paragraph("差旅与费用报销管理制度", title),
            Paragraph("文档编号：HR-004　版本：V1.0　生效日期：2026年10月1日", meta),
            Paragraph("制度归口：财务部、人力资源部", meta), Paragraph(NOTICE, meta), Spacer(1, 4 * mm)]
    rows = [["住宿城市", "住宿上限（元/人/晚）"], ["北京、上海、深圳", "600"], ["其他中国境内城市", "450"]]
    table = Table([[Paragraph(cell, body) for cell in row] for row in rows], colWidths=[85*mm, 79*mm])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5F1EC")),
                               ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCD9D2")),
                               ("TOPPADDING", (0, 0), (-1, -1), 8),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    for index, (name, paragraphs) in enumerate(TRAVEL):
        flow.append(Paragraph(name, heading))
        flow.extend(Paragraph(text, body) for text in paragraphs)
        if index == 1:
            flow.append(KeepTogether([Paragraph("住宿标准速查", meta), table]))

    def page_footer(canvas, document):
        canvas.setStrokeColor(colors.HexColor("#CCD9D2"))
        canvas.line(22*mm, 19*mm, 188*mm, 19*mm)
        canvas.setFont("Chinese", 8)
        canvas.setFillColor(colors.HexColor("#53615D"))
        canvas.drawString(22*mm, 14*mm, "虚构测试资料 | HR-004 | V1.0")
        canvas.drawRightString(188*mm, 14*mm, f"第 {document.page} 页")

    pdf = SimpleDocTemplate(str(OUTPUT / "差旅与费用报销管理制度.pdf"),
                            title="差旅与费用报销管理制度", author=COMPANY + "（虚构测试）",
                            pagesize=(210*mm, 297*mm), leftMargin=22*mm, rightMargin=22*mm,
                            topMargin=20*mm, bottomMargin=25*mm)
    pdf.build(flow, onFirstPage=page_footer, onLaterPages=page_footer)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/simhei.ttf"),
                        help="支持中文的本地 TrueType 字体，用于 PDF 字体嵌入")
    args = parser.parse_args()
    if not args.font.is_file():
        parser.error("找不到中文字体，请通过 --font 指定支持中文的 TTF 字体")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    create_word()
    create_pdf(args.font)
    print("已生成招聘制度 DOCX 和差旅制度 PDF：", OUTPUT)
